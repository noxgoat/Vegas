"""Exportação dos resultados (CSV e XLSX) com proteção contra injeção de fórmulas."""
from __future__ import annotations

import csv
import io
import json
import sqlite3

from .query_parser import MODE_LABEL
from .search import NAO_LOCALIZADO

COLUMNS = [
    ("escopo_da_exportacao", "Escopo da exportação"), ("termo", "Termo"), ("modo", "Modo de correspondência"),
    ("ocorrencia", "Ocorrência nº"), ("trecho", "Trecho original"), ("texto_correspondente", "Texto correspondente"),
    ("pagina_pdf", "Página do PDF"), ("posicao_inicio", "Posição inicial"), ("posicao_fim", "Posição final"),
    ("titulo", "Título"), ("data_disponibilizacao", "Data de disponibilização"),
    ("data_publicacao", "Data de publicação"), ("data_decisao", "Data da decisão"), ("tribunal", "Tribunal"),
    ("orgao", "Órgão julgador / vara"), ("tipo_ato", "Tipo de ato (informado pela fonte)"),
    ("processo", "Número do processo"), ("magistrado", "Magistrado / relator"),
    ("url_publicacao", "URL da publicação"), ("url_original", "URL do documento original"),
    ("fonte", "Fonte"), ("ultima_coleta", "Última coleta"), ("estado_extracao", "Estado da extração"),
    ("estado_verificacao", "Estado da verificação"), ("documento_id", "ID do documento"),
    ("versao_pesquisada", "Versão pesquisada"),
]
DANGEROUS = ("=", "+", "-", "@", "\t", "\r", "\n")


def safe_cell(v) -> str:
    """Neutraliza conteúdo que planilhas interpretariam como fórmula."""
    if v is None:
        return ""
    s = str(v)
    if s and s[0] in DANGEROUS:
        return "'" + s
    return s


def iter_rows(conn: sqlite3.Connection, run_id: int, term_index: int | None = None,
              source_id: int | None = None):
    run = conn.execute("SELECT * FROM search_runs WHERE id = ?", (run_id,)).fetchone()
    if run is None:
        raise KeyError("Pesquisa não encontrada")
    filters = []
    if term_index is not None:
        terms = {t["index"]: t["raw"] for t in json.loads(run["terms_json"])}
        filters.append(f"termo={terms.get(term_index, term_index)!r}")
    if source_id:
        name = conn.execute("SELECT name FROM sources WHERE id = ?", (source_id,)).fetchone()
        filters.append(f"fonte={name[0] if name else source_id}")
    scope = ("Seleção filtrada: " + ", ".join(filters)) if filters else "Todos os resultados da pesquisa"
    where = "o.search_run_id = ?"
    args: list = [run_id]
    if term_index is not None:
        where += " AND o.term_index = ?"
        args.append(term_index)
    if source_id:
        where += " AND d.source_id = ?"
        args.append(source_id)
    q = (f"SELECT o.*, d.*, s.name AS source_name, o.document_version AS searched_version FROM occurrences o "
         f"JOIN documents d ON d.id = o.document_id JOIN sources s ON s.id = d.source_id "
         f"WHERE {where} ORDER BY o.term_index, d.id, o.ordinal")
    for r in conn.execute(q, args):
        def v(x):
            return x if x not in (None, "") else NAO_LOCALIZADO
        verified = r["extraction_status"] == "completo" and r["text_scope"] == "integral"
        yield {
            "escopo_da_exportacao": scope, "termo": r["term"], "modo": MODE_LABEL.get(r["match_mode"], r["match_mode"]),
            "ocorrencia": r["ordinal"],
            "trecho": " ".join((r["snippet_before"] + "[" + r["match_text"] + "]" + r["snippet_after"]).split()),
            "texto_correspondente": r["match_text"], "pagina_pdf": r["page"] or "",
            "posicao_inicio": r["start_pos"], "posicao_fim": r["end_pos"],
            "titulo": v(r["title"]), "data_disponibilizacao": v(r["date_available"]),
            "data_publicacao": v(r["date_publication"]), "data_decisao": v(r["date_decision"]),
            "tribunal": v(r["court"]), "orgao": v(r["organ"]), "tipo_ato": v(r["act_type"]),
            "processo": v(r["process_number"]), "magistrado": v(r["judge"]),
            "url_publicacao": v(r["url_publication"]), "url_original": v(r["url_original"]),
            "fonte": r["source_name"], "ultima_coleta": r["last_collected_at"],
            "estado_extracao": r["extraction_status"],
            "estado_verificacao": "Integralmente verificado" if verified else "Parcial",
            "documento_id": r["document_id"], "versao_pesquisada": r["searched_version"],
        }


def to_csv(conn, run_id, term_index=None, source_id=None) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";", quoting=csv.QUOTE_ALL, lineterminator="\r\n")
    w.writerow([label for _, label in COLUMNS])
    for row in iter_rows(conn, run_id, term_index, source_id):
        w.writerow([safe_cell(row[k]) for k, _ in COLUMNS])
    return ("﻿" + buf.getvalue()).encode("utf-8")   # BOM para o Excel reconhecer UTF-8


def to_xlsx(conn, run_id, term_index=None, source_id=None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Resultados"
    ws.append([label for _, label in COLUMNS])
    for row in iter_rows(conn, run_id, term_index, source_id):
        ws.append([safe_cell(row[k]) for k, _ in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
